"""
ASSET REGISTER (Tujuan-1 TA) — ubah hasil digitisasi P&ID jadi daftar aset yang rapi,
lalu tulis ke Word (.docx).

Isi register:
  1. Equipment    — dengan TAG dibaca dari gambar (OCR di sekitar kotak equipment).
  2. Instrument   — tag ISA (fungsi + nomor loop) dibaca dari bubble.
  3. Valve        — rekap per jenis + daftar posisi.
  4. Piping (line register) — line number, unit, size, process fluid, piping class,
     material, corrosion system & circuit (Fitur 2).
  5. Connection point (spec break) bila ada.

Catatan metodologi (dipakai di laporan): TAG yang terbaca adalah hasil OCR, jadi
register menandai tag yang TIDAK terbaca sebagai "(not read)" alih-alih mengarang —
angka yang tidak bisa dipertanggungjawabkan tidak boleh masuk dokumen aset.
"""
from __future__ import annotations

import os
import re
from collections import Counter, OrderedDict
from datetime import date

from .systemize import circuitize, material_of
from .subtype import isa_describe, sym_class

# tag equipment: 605-V-201, 610-E-101A, 50-K-102
_EQ_TAG = re.compile(r"\d{2,3}\s*-\s*[A-Z]{1,3}\s*-\s*\d{2,4}\s*[A-Z]?")
# tag instrument ISA: huruf fungsi (2-4) + nomor loop
_ISA_TAG = re.compile(r"^([A-Z]{2,4})[\s\-]*([0-9][0-9A-Z\-]*)?$")
# nomor loop: 2-4 angka + akhiran huruf opsional. Lebih longgar dari ini dan angka lain di
# sekitar bubble (setpoint, elevasi, ukuran) ikut tertangkap sebagai 'loop'.
_LOOP_RE = re.compile(r"\d{2,4}[A-Z]{0,2}")
_EQ_KW = ("PUMP", "TANK", "DRUM", "VESSEL", "EXCHANGER", "COOLER", "HEATER", "COMPRESSOR",
          "FILTER", "COLUMN", "TOWER", "SEPARATOR", "SCRUBBER", "RECEIVER", "REACTOR", "POT",
          "SKID", "PACKAGE", "UNIT", "CONSOLE", "COALESCER", "LAUNCHER", "BLOWER", "CATCHER",
          "KNOCKOUT", "REBOILER", "CONDENSER", "STRIPPER", "ACCUMULATOR", "SILENCER")

NOT_READ = "(not read)"


# --------------------------------------------------------------------- OCR ------
def _ocr_boxes(img_bgr, x1, y1, x2, y2, pad=6, angles=(0,)):
    """OCR sebuah region -> [(teks, cx, cy)] dalam KOORDINAT GAMBAR ASLI.

    Teks TIDAK dinormalisasi ke format piping ID (tag equipment/instrument punya format
    sendiri). Posisi dikembalikan supaya pemanggil bisa memilih tag TERDEKAT — di P&ID
    tag equipment ditulis di samping/atas simbolnya, dan tetangga bisa ikut terbaca.
    Angle 0 saja secara default: tag equipment hampir selalu horizontal, dan multi-sudut
    melipatgandakan waktu OCR."""
    import cv2
    from .piping_id import get_rapid

    H, W = img_bgr.shape[:2]
    xa, ya = max(0, int(min(x1, x2)) - pad), max(0, int(min(y1, y2)) - pad)
    xb, yb = min(W, int(max(x1, x2)) + pad), min(H, int(max(y1, y2)) + pad)
    if xb - xa < 4 or yb - ya < 4:
        return []
    crop = img_bgr[ya:yb, xa:xb]
    if crop.size == 0:
        return []
    f = 1.0
    if crop.shape[0] < 48:                       # teks kecil -> upscale dulu
        f = 48 / crop.shape[0]
        crop = cv2.resize(crop, None, fx=f, fy=f, interpolation=cv2.INTER_CUBIC)
    eng = get_rapid()
    out = []
    for ang in angles:
        rot = crop if ang == 0 else cv2.rotate(
            crop, {90: cv2.ROTATE_90_CLOCKWISE, 270: cv2.ROTATE_90_COUNTERCLOCKWISE}[ang])
        try:
            res, _ = eng(cv2.cvtColor(rot, cv2.COLOR_BGR2RGB))
        except Exception:
            res = None
        for b, t, _s in (res or []):
            t = (t or "").strip().upper()
            if not t:
                continue
            if ang == 0:
                xs = [p[0] for p in b]; ys = [p[1] for p in b]
                cx, cy = xa + sum(xs) / len(xs) / f, ya + sum(ys) / len(ys) / f
            else:                                # sudut lain: pakai pusat region saja
                cx, cy = (xa + xb) / 2, (ya + yb) / 2
            out.append((t, cx, cy))
    return out


def _clean_tag(t: str) -> str:
    return re.sub(r"\s*-\s*", "-", t.strip().upper()).replace(" ", "")


# Huruf tengah tag equipment -> jenis. Konvensi umum industri; ASUMSI yang bisa diganti
# kalau perusahaan punya daftar resmi sendiri.
_TYPE_FROM_TAG = {"V": "Vessel", "D": "Drum", "T": "Tank", "C": "Column", "E": "Heat exchanger",
                  "K": "Compressor", "P": "Pump", "F": "Filter / coalescer", "R": "Reactor",
                  "H": "Heater", "S": "Separator", "B": "Blower", "X": "Package / skid"}


def _describe_tag(tag: str) -> str:
    m = re.match(r"^\d{2,3}-([A-Z]{1,3})-", tag or "")
    return _TYPE_FROM_TAG.get(m.group(1)[:1], "Equipment") if m else ""


def read_equipment_tag(img_bgr, sym) -> tuple[str, str]:
    """(tag, description) sebuah equipment.

    Tag equipment di P&ID ditulis DI SAMPING atau di atas simbolnya (mis. `610-V-202A`
    di kanan vessel), bukan di dalamnya — jadi yang di-OCR adalah kotak yang DIPERLUAS,
    lalu dipilih tag yang PALING DEKAT ke kotak supaya tag tetangga tidak tercuri."""
    x1, y1, x2, y2 = sym["x1"], sym["y1"], sym["x2"], sym["y2"]
    w, h = max(20.0, x2 - x1), max(20.0, y2 - y1)
    ex, ey = max(90.0, 0.95 * w), max(70.0, 0.75 * h)
    cx0, cy0 = (x1 + x2) / 2, (y1 + y2) / 2

    found = _ocr_boxes(img_bgr, x1 - ex, y1 - ey, x2 + ex, y2 + ey, pad=0)
    cands = []
    for t, cx, cy in found:
        m = _EQ_TAG.search(t.replace(" ", ""))
        if m:
            cands.append((abs(cx - cx0) + abs(cy - cy0), _clean_tag(m.group(0))))
    tag = min(cands)[1] if cands else ""

    desc = _describe_tag(tag)
    if not desc:                                  # tak ada tag -> coba nama tertulis
        words = {w2 for t, _c, _y in found for w2 in re.split(r"[^A-Z]+", t) if w2}
        kw = next((k for k in _EQ_KW if k in words), None)   # kata UTUH, bukan substring
        if kw:
            cand = [t for t, _c, _y in found if kw in re.split(r"[^A-Z]+", t)]
            desc = max(cand, key=len).title() if cand else kw.title()
    return tag, desc


def _bubble_read(img_bgr, sym, pad):
    got = _ocr_boxes(img_bgr, sym["x1"], sym["y1"], sym["x2"], sym["y2"], pad=pad, angles=(0,))
    got.sort(key=lambda g: g[2])                  # atas -> bawah: fungsi dulu, nomor loop di bawah
    func, loop = "", ""
    for t, _x, _y in got:
        c = _clean_tag(t)
        m = _ISA_TAG.match(c)
        if m:
            if not func:
                func = m.group(1)
            if m.group(2) and not loop and _LOOP_RE.fullmatch(m.group(2)):
                loop = m.group(2)
        elif not loop and _LOOP_RE.fullmatch(c):
            loop = c
    return func, loop


def read_instrument_tag(img_bgr, sym) -> tuple[str, str]:
    """(fungsi ISA, nomor loop) sebuah instrument dari bubble-nya.

    Bubble instrumen selalu horizontal -> OCR sudut 0 saja: 3x lebih cepat dan lebih sedikit
    pembacaan sampah. Kalau nomor loop keluar >=4 digit, biasanya UJUNG PANAH leader yang
    menempel di tepi bubble ikut terbaca sebagai angka (mis. 'PI 129' -> '7129'); dicoba
    ulang dengan crop lebih rapat, dan hasilnya HANYA diterima bila merupakan sufiks dari
    bacaan pertama — supaya perbaikan ini tidak pernah mengarang angka baru."""
    func, loop = _bubble_read(img_bgr, sym, pad=3)
    if len(re.sub(r"\D", "", loop)) >= 4:
        inset = -max(2, int(min(sym["x2"] - sym["x1"], sym["y2"] - sym["y1"]) * 0.07))
        f2, l2 = _bubble_read(img_bgr, sym, pad=inset)
        if l2 and l2 != loop and loop.endswith(l2):
            loop = l2
            func = func or f2
    if not func:
        func = (sym.get("subtype") or "").strip().upper()
    return func, loop


# ------------------------------------------------------------------ register ----
def _pos(sym):
    return int((sym["x1"] + sym["x2"]) / 2), int((sym["y1"] + sym["y2"]) / 2)


def build_register(result, img_bgr=None, read_tags=True, progress=None) -> dict:
    """Susun seluruh daftar aset dari `result`. Bila `img_bgr` diberikan dan
    read_tags=True, tag equipment & instrument DIBACA dari gambar (OCR) lalu
    di-cache ke symbol supaya ekspor berikutnya cepat."""
    syms = result.get("symbols") or []
    pids = result.get("piping_ids") or []
    eqs = [s for s in syms if s.get("coarse") == "equipment"]
    ins = [s for s in syms if s.get("coarse") == "instrument"]
    vls = [s for s in syms if s.get("coarse") == "valve"]

    # ---- equipment ----
    equipment = []
    for i, s in enumerate(sorted(eqs, key=lambda s: (s["y1"], s["x1"])), 1):
        tag, desc = s.get("tag", ""), s.get("desc", "")
        manual = (s.get("subtype") or "").strip()
        if not tag and not desc and img_bgr is not None and read_tags:
            if progress:
                progress(f"reading equipment tag {i}/{len(eqs)}…")
            tag, desc = read_equipment_tag(img_bgr, s)
            s["tag"], s["desc"] = tag, desc            # cache
        if manual and not desc and not _EQ_TAG.search(manual):
            desc = manual                              # nama yang diketik user di GUI
        if manual and not tag and _EQ_TAG.search(manual):
            tag = _clean_tag(manual)
        x, y = _pos(s)
        equipment.append({"no": i, "tag": tag or NOT_READ, "desc": desc or "Equipment",
                          "source": "manual" if manual else "detected", "x": x, "y": y})

    # Prefix unit tag instrumen (mis. '610' pada 610-PI-129) ditulis di luar bubble. Diambil
    # dari UNIT line number gambar ini, bukan OCR terpisah: nilainya sama dan sudah terbaca.
    units = Counter(p.get("unit") for p in pids if p.get("unit"))
    unit = units.most_common(1)[0][0] if units else ""

    # ---- instrument ----
    instrument = []
    for i, s in enumerate(sorted(ins, key=lambda s: (s["y1"], s["x1"])), 1):
        func, loop = s.get("isa_func", ""), s.get("isa_loop", "")
        if not func and img_bgr is not None and read_tags:
            if progress and i % 10 == 1:
                progress(f"reading instrument tag {i}/{len(ins)}…")
            func, loop = read_instrument_tag(img_bgr, s)
            s["isa_func"], s["isa_loop"] = func, loop
        if not func:
            func = (s.get("subtype") or "").strip().upper()
        tag = "-".join(t for t in (unit, func, loop) if t) if func else NOT_READ
        x, y = _pos(s)
        instrument.append({"no": i, "tag": tag, "func": func or NOT_READ,
                           "desc": isa_describe(func) if func else "Instrument",
                           "x": x, "y": y})

    # ---- valve ----
    vcount = Counter(sym_class(s) for s in vls)
    valve = [{"no": i, "type": k, "count": v}
             for i, (k, v) in enumerate(sorted(vcount.items(), key=lambda kv: -kv[1]), 1)]
    valve_items = [{"no": i, "type": sym_class(s), "x": _pos(s)[0], "y": _pos(s)[1]}
                   for i, s in enumerate(sorted(vls, key=lambda s: (s["y1"], s["x1"])), 1)]

    # ---- piping (line register) + corrosion system/circuit ----
    systems = circuitize(result)
    piping, sys_rows, cir_rows = [], [], []
    n = 0
    for s in systems:
        classes = sorted({cl for c in s["circuits"] for cl in c["classes"]})
        sys_rows.append({"code": f"#{s['index']:02d}", "fluid": s["fluid"],
                         "n_line": len(s["pid_idxs"]), "n_pipe": s["n_pipes"],
                         "n_circuit": len(s["circuits"]), "classes": ", ".join(classes),
                         "color": tuple(s["color"])})
        for c in s["circuits"]:
            cir_rows.append({"code": c["code"], "fluid": s["fluid"],
                             "material": c.get("material") or "—",
                             "classes": ", ".join(sorted(c["classes"])),
                             "n_line": len(c["pid_idxs"]), "color": tuple(c["color"])})
            for pi in c["pid_idxs"]:
                p = pids[pi]; n += 1
                piping.append({"no": n, "line": p.get("pid", ""), "unit": p.get("unit", ""),
                               "size": p.get("size", ""), "fluid": s["fluid"],
                               "pclass": p.get("pclass", ""),
                               "material": material_of(p.get("pclass")) or "—",
                               "system": f"#{s['index']:02d}", "circuit": c["code"],
                               "state": p.get("state", ""), "color": tuple(s["color"])})
    listed = {r["line"] for r in piping}
    for p in pids:                                    # line tanpa fluid -> tetap terdaftar
        if p.get("pid") and p["pid"] not in listed:
            n += 1
            piping.append({"no": n, "line": p["pid"], "unit": p.get("unit", ""),
                           "size": p.get("size", ""), "fluid": p.get("fluid", "") or "—",
                           "pclass": p.get("pclass", ""),
                           "material": material_of(p.get("pclass")) or "—",
                           "system": "—", "circuit": "—", "state": p.get("state", ""),
                           "color": None})

    cps = []
    for i, cp in enumerate(result.get("conn_points") or [], 1):
        cps.append({"no": i, "codes": " | ".join(cp.get("codes") or []),
                    "attached": "yes" if cp.get("run_idx") is not None else "no",
                    "x": int(cp.get("x", 0)), "y": int(cp.get("y", 0))})

    return {
        "equipment": equipment, "instrument": instrument,
        "valve_summary": valve, "valve_items": valve_items,
        "piping": piping, "systems": sys_rows, "circuits": cir_rows,
        "conn_points": cps,
        "summary": OrderedDict([
            ("Equipment", len(equipment)),
            ("Instrument", len(instrument)),
            ("Valve", len(valve_items)),
            ("Piping line (line number)", len(piping)),
            ("Corrosion system", len(sys_rows)),
            ("Corrosion circuit", len(cir_rows)),
            ("Connection point (spec break)", len(cps)),
        ]),
    }


# ------------------------------------------------------------------- Word -------
def _shade(cell, hexcolor):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")                  # 'clear', bukan 'solid' (solid = hitam)
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)


def _table(doc, headers, widths_cm, rows, style="Table Grid", font_pt=8.5,
           shade_col=None, header_fill="1F4E96"):
    """Tabel dengan lebar kolom EKSPLISIT di tiap sel (kalau tidak, Word/Docs mengabaikannya).
    shade_col = {index_kolom: fungsi(row)->'RRGGBB' atau None} untuk mewarnai sel grup."""
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.table import WD_TABLE_ALIGNMENT

    t = doc.add_table(rows=1, cols=len(headers))
    t.style = style
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr = t.rows[0].cells
    for j, (h, w) in enumerate(zip(headers, widths_cm)):
        hdr[j].width = Cm(w)
        p = hdr[j].paragraphs[0]; p.text = ""
        r = p.add_run(h); r.bold = True; r.font.size = Pt(font_pt)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(hdr[j], header_fill)
    for row in rows:
        cells = t.add_row().cells
        for j, (val, w) in enumerate(zip(row, widths_cm)):
            cells[j].width = Cm(w)
            p = cells[j].paragraphs[0]; p.text = ""
            r = p.add_run("" if val is None else str(val)); r.font.size = Pt(font_pt)
        if shade_col:
            for j, fn in shade_col.items():
                col = fn(row)
                if col:
                    _shade(cells[j], col)
    return t


def _h(doc, text, level=1):
    from docx.shared import Pt, RGBColor
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x96)
        r.font.size = Pt(15 if level == 1 else 12)
    return p


def _note(doc, text):
    from docx.shared import Pt, RGBColor
    p = doc.add_paragraph()
    r = p.add_run(text); r.italic = True; r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x7A, 0x87, 0x9A)
    return p


def _hex(rgb):
    return None if not rgb else "%02X%02X%02X" % tuple(int(c) for c in rgb)


def export_asset_register_docx(result, path, drawing_name="", img_bgr=None,
                               read_tags=True, progress=None, reg=None) -> dict:
    """Tulis ASSET REGISTER ke .docx. Return dict ringkasan jumlah aset."""
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    reg = reg or build_register(result, img_bgr=img_bgr, read_tags=read_tags, progress=progress)
    if progress:
        progress("writing Word document…")

    doc = Document()
    sec = doc.sections[0]                          # A4 landscape: line register itu lebar
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Cm(1.5))
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(9)

    ttl = doc.add_heading("P&ID ASSET REGISTER", level=0)
    ttl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = sub.add_run(drawing_name or os.path.basename(result.get("image_path") or ""))
    rr.bold = True; rr.font.size = Pt(11)
    _note(doc, f"Generated automatically from the P&ID digitization result  ·  "
               f"{date.today().isoformat()}  ·  grouping per API RP 970 §5.6").alignment = \
        WD_ALIGN_PARAGRAPH.CENTER

    # ---- 1. summary ----
    _h(doc, "1.  Summary", 1)
    _table(doc, ["Asset category", "Quantity"], [10.0, 4.0],
           [[k, v] for k, v in reg["summary"].items()], font_pt=9.5)

    # ---- 2. equipment ----
    _h(doc, "2.  Equipment list", 1)
    if reg["equipment"]:
        _table(doc, ["No", "Tag", "Description", "Source", "X", "Y"],
               [1.2, 4.5, 9.0, 2.6, 2.2, 2.2],
               [[e["no"], e["tag"], e["desc"], e["source"], e["x"], e["y"]]
                for e in reg["equipment"]])
        _note(doc, "Tag is read by OCR from the drawing. Entries shown as "
                   f"“{NOT_READ}” were detected as equipment but carry no readable "
                   "tag — they must be filled in by the engineer rather than guessed.")
    else:
        _note(doc, "No equipment detected on this drawing.")

    # ---- 3. instrument ----
    _h(doc, "3.  Instrument list", 1)
    if reg["instrument"]:
        func_count = Counter(i["func"] for i in reg["instrument"])
        _table(doc, ["ISA function", "Description", "Quantity"], [3.5, 9.0, 3.0],
               [[k, isa_describe(k) if k != NOT_READ else "—", v]
                for k, v in sorted(func_count.items(), key=lambda kv: -kv[1])], font_pt=9)
        doc.add_paragraph()
        _table(doc, ["No", "Tag", "ISA function", "Description", "X", "Y"],
               [1.2, 4.0, 3.0, 8.0, 2.2, 2.2],
               [[i["no"], i["tag"], i["func"], i["desc"], i["x"], i["y"]]
                for i in reg["instrument"]])
        _note(doc, "ISA function follows ISA-5.1 letter coding (first letter = measured "
                   "variable, following letters = function). The unit prefix in the tag is "
                   "taken from the unit token of this drawing's line numbers, not read "
                   "separately from each bubble.")
    else:
        _note(doc, "No instrument detected on this drawing.")

    # ---- 4. valve ----
    _h(doc, "4.  Valve list", 1)
    if reg["valve_items"]:
        _table(doc, ["No", "Valve type", "Quantity"], [1.2, 9.0, 3.0],
               [[v["no"], v["type"], v["count"]] for v in reg["valve_summary"]], font_pt=9)
        doc.add_paragraph()
        _table(doc, ["No", "Valve type", "X", "Y"], [1.2, 9.0, 2.4, 2.4],
               [[v["no"], v["type"], v["x"], v["y"]] for v in reg["valve_items"]])
    else:
        _note(doc, "No valve detected on this drawing.")

    # ---- 5. piping ----
    doc.add_page_break()
    _h(doc, "5.  Piping list (line register)", 1)
    if reg["piping"]:
        _table(doc, ["No", "Line number", "Unit", "Size", "Process fluid", "Piping class",
                     "Material", "Corrosion system", "Circuit", "Pipe"],
               [1.1, 7.2, 1.6, 1.6, 2.6, 2.4, 2.2, 3.2, 2.2, 2.3],
               [[p["no"], p["line"], p["unit"], p["size"], p["fluid"], p["pclass"],
                 p["material"], p["system"], p["circuit"], p["state"]] for p in reg["piping"]],
               shade_col={7: lambda r: None})
        # warnai kolom corrosion system sesuai warna marking
        tbl = doc.tables[-1]
        for i, p in enumerate(reg["piping"], 1):
            if p["color"]:
                _shade(tbl.rows[i].cells[7], _hex(p["color"]))
        _note(doc, "“Pipe” states how the line number was linked to its pipe: "
                   "attached = directly below/beside the text, leader = via a leader line, "
                   "propagated = inherited through connectivity, none = not linked yet.")
    else:
        _note(doc, "No piping ID read on this drawing.")

    # ---- 6. corrosion system & circuit ----
    _h(doc, "6.  Corrosion system & circuit (API RP 970)", 1)
    if reg["systems"]:
        _table(doc, ["System", "Process fluid", "Lines", "Traced pipes", "Circuits",
                     "Piping classes involved"], [2.2, 3.0, 2.0, 2.8, 2.0, 10.0],
               [[s["code"], s["fluid"], s["n_line"], s["n_pipe"], s["n_circuit"], s["classes"]]
                for s in reg["systems"]], font_pt=9)
        t = doc.tables[-1]
        for i, s in enumerate(reg["systems"], 1):
            _shade(t.rows[i].cells[0], _hex(s["color"]))
        doc.add_paragraph()
        _table(doc, ["Circuit", "System fluid", "Material", "Piping classes", "Lines"],
               [2.4, 3.0, 2.4, 9.0, 2.0],
               [[c["code"], c["fluid"], c["material"], c["classes"], c["n_line"]]
                for c in reg["circuits"]], font_pt=9)
        t = doc.tables[-1]
        for i, c in enumerate(reg["circuits"], 1):
            _shade(t.rows[i].cells[0], _hex(c["color"]))
        _note(doc, "Systemization groups pipes by process fluid (API RP 970 §5.6, "
                   "“process stream composition”). Circuitization then splits each "
                   "system by the material implied by the piping class.")
    else:
        _note(doc, "No corrosion system could be formed (no process fluid read).")

    # ---- 7. connection points ----
    if reg["conn_points"]:
        _h(doc, "7.  Connection point (piping-class spec break)", 1)
        _table(doc, ["No", "Class change", "Attached to a pipe", "X", "Y"],
               [1.2, 6.0, 4.2, 2.4, 2.4],
               [[c["no"], c["codes"], c["attached"], c["x"], c["y"]]
                for c in reg["conn_points"]], font_pt=9)
        _note(doc, "A connection point is a piping-class change written explicitly on the "
                   "P&ID by its designer — a natural candidate for a circuit boundary.")

    # ---- method note ----
    doc.add_page_break()
    _h(doc, "Method note", 1)
    for line in [
        "Symbol detection (equipment, valve, instrument) uses YOLO11 object detectors; "
        "valve type is refined by a dedicated shape classifier.",
        "Piping ID text is located and read with a generic OCR engine plus a configurable "
        "pattern parser — no per-company annotation is required.",
        "Equipment and instrument tags in this register are read by OCR from the drawing "
        "itself; anything unreadable is reported as “not read” and is never invented.",
        "Line tracing, grouping (systemization / circuitization) and marking are deterministic "
        "rule-based steps, not machine learning, so every grouping decision can be traced back "
        "to the line number it came from.",
        "This register is a computed draft. It is intended to be verified by a corrosion "
        "engineer before use in a Corrosion Control Document.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    doc.save(path)
    return dict(reg["summary"])
